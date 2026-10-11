"""Live Excel mirror of villager login accounts.

Rewrites exports/login_accounts_management.xlsx so it always reflects the real
citizen (villager) accounts and their profile / login state. Mirrors the
notice-board Excel pattern: the sync is best-effort and never breaks the API.

Called automatically whenever a villager account is created (registration) or
logs in, and exposed for download from the Panchayat Desk.
"""

import os
from datetime import datetime, timezone
from pathlib import Path

from app.supabase_client import supabase

VILLAGER_XLSX = Path(
    os.getenv(
        "LOGIN_ACCOUNTS_XLSX",
        str(Path(__file__).resolve().parents[2] / "exports" / "login_accounts_management.xlsx"),
    )
)

HEADERS = [
    "User ID", "Full Name", "Mobile Number", "Email", "Account Status",
    "Village", "Ward", "House No", "Address", "Date of Birth", "Gender",
    "Occupation", "Verification", "Registered On", "Last Login",
]


def _fetch_all(table: str, columns: str, **eq) -> list:
    """Page through a table (PostgREST caps a single select at 1000 rows)."""
    rows, offset = [], 0
    while True:
        q = supabase.table(table).select(columns)
        for col, val in eq.items():
            q = q.eq(col, val)
        page = q.order("created_at", desc=False).range(offset, offset + 999).execute().data or []
        rows.extend(page)
        if len(page) < 1000:
            break
        offset += 1000
    return rows


def _lookup(table: str, key: str, label: str) -> dict:
    rows, offset = [], 0
    while True:
        page = (
            supabase.table(table)
            .select(f"{key},{label}")
            .order(key)
            .range(offset, offset + 999)
            .execute()
            .data
            or []
        )
        rows.extend(page)
        if len(page) < 1000:
            break
        offset += 1000
    return {r[key]: r.get(label) for r in rows}


def sync_villagers_excel() -> str:
    """Rewrite exports/login_accounts_management.xlsx. Never raises."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        users = _fetch_all(
            "users",
            "user_id,full_name,mobile_number,email,role,account_status,last_login,created_at",
            role="villager",
        )
        profiles = {
            p["user_id"]: p
            for p in _fetch_all(
                "villagers",
                "user_id,village_id,ward_id,house_number,address,date_of_birth,gender,occupation,verification_status",
            )
        }
        villages = _lookup("villages", "village_id", "village_name")
        wards = _lookup("wards", "ward_id", "ward_name")

        wb = Workbook()
        ws = wb.active
        ws.title = "Login Accounts"
        ws.append(HEADERS)

        head_fill = PatternFill("solid", fgColor="1F4E44")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        synced_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
        for u in users:
            p = profiles.get(u["user_id"], {})
            village = villages.get(p.get("village_id"), "")
            ward = p.get("ward_id")
            ward_label = ""
            if ward:
                ward_label = wards.get(ward) or ""
            ws.append([
                u.get("user_id"),
                u.get("full_name"),
                u.get("mobile_number"),
                u.get("email") or "",
                u.get("account_status") or "",
                village,
                ward_label,
                p.get("house_number") or "",
                p.get("address") or "",
                (p.get("date_of_birth") or "")[:10],
                p.get("gender") or "",
                p.get("occupation") or "",
                p.get("verification_status") or "",
                (u.get("created_at") or "")[:19].replace("T", " "),
                (u.get("last_login") or "")[:19].replace("T", " "),
            ])

        # Footer stamp so anyone reading the sheet knows when it was last synced.
        ws.append([])
        ws.append([f"Synced at {synced_at} UTC  |  {len(users)} villager account(s)"])

        for col, width in {
            "A": 38, "B": 24, "C": 14, "D": 26, "E": 14, "F": 18, "G": 12,
            "H": 12, "I": 30, "J": 14, "K": 10, "L": 18, "M": 14, "N": 20, "O": 20,
        }.items():
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"

        VILLAGER_XLSX.parent.mkdir(parents=True, exist_ok=True)
        wb.save(VILLAGER_XLSX)
        return str(VILLAGER_XLSX)
    except Exception:
        return ""

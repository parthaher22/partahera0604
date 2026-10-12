from datetime import datetime, timedelta, timezone
import os
import re
import time
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Request
from pydantic import BaseModel
from typing import Optional

from app.audit_util import log_audit
from app.config import KUMBHARI_VILLAGE_ID
from app.supabase_client import supabase
from app.security import get_current_user, require_roles, CurrentUser

router = APIRouter(prefix="/api/announcements", tags=["announcements"])

VALID_STATUSES = {"Draft", "Published", "Archived"}

# The citizen notice board keeps only the newest N notices (older ones are
# cleared from the board but stay in the desk's archive).
CITIZEN_NOTICE_LIMIT = int(os.getenv("CITIZEN_NOTICE_LIMIT", "6"))
STAFF_NOTICE_LIMIT = 500

# Excel file that mirrors the Digital Notice Board in real time.
NOTICE_BOARD_XLSX = Path(
    os.getenv(
        "NOTICE_BOARD_XLSX",
        str(Path(__file__).resolve().parents[3] / "exports" / "notice_board.xlsx"),
    )
)


class AnnouncementCreate(BaseModel):
    village_id: Optional[str] = None
    title: str
    description: str
    category: str = "General Notice"
    priority: str = "Medium"
    event_date: Optional[str] = None
    expiry_date: Optional[str] = None
    attachment_url: Optional[str] = None


class AnnouncementUpdate(BaseModel):
    village_id: Optional[str] = None
    title: Optional[str] = None
    description: Optional[str] = None
    category: Optional[str] = None
    priority: Optional[str] = None
    status: Optional[str] = None
    event_date: Optional[str] = None
    expiry_date: Optional[str] = None
    attachment_url: Optional[str] = None
    # Scheduled publishing: when status becomes Published and this timestamp is
    # in the future, citizens only see the notice once published_at is reached.
    publish_at: Optional[str] = None


LANG_TAG_RE = re.compile(r"^\s*---(mr|hi)---\s*$")


def split_langs(text: str) -> dict:
    """Split a notice description into {en, mr, hi} sections.

    Multilingual notices are stored in one text column using language tags:
        <english text>
        ---mr---
        <marathi text>
        ---hi---
        <hindi text>
    """
    out = {"en": "", "mr": "", "hi": ""}
    current = "en"
    for line in (text or "").splitlines():
        m = LANG_TAG_RE.match(line)
        if m:
            current = m.group(1)
            continue
        out[current] += line + "\n"
    return {k: v.strip() for k, v in out.items()}


def _sync_notice_excel() -> str:
    """Rewrite exports/notice_board.xlsx so the Digital Notice Board Excel
    always mirrors the live data. Failures never break the API."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        rows = (
            supabase.table("announcements")
            .select("*")
            .eq("village_id", KUMBHARI_VILLAGE_ID)
            .order("created_at", desc=True)
            .limit(1000)
            .execute()
            .data
            or []
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Notice Board"
        headers = [
            "Title", "Category", "Priority", "Status", "Published On", "Expiry",
            "Description (English)", "Description (Marathi)", "Description (Hindi)",
            "Attachment", "Synced At",
        ]
        ws.append(headers)
        head_fill = PatternFill("solid", fgColor="1F4E44")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill

        synced_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
        for a in rows:
            langs = split_langs(a.get("description"))
            ws.append([
                a.get("title"),
                a.get("category"),
                a.get("priority"),
                a.get("status"),
                (a.get("published_at") or "")[:19].replace("T", " "),
                a.get("expiry_date") or "",
                langs["en"],
                langs["mr"],
                langs["hi"],
                a.get("attachment_url") or "",
                synced_at,
            ])

        for col, width in {
            "A": 46, "B": 18, "C": 11, "D": 12, "E": 20, "F": 13,
            "G": 70, "H": 70, "I": 70, "J": 30, "K": 20,
        }.items():
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"

        NOTICE_BOARD_XLSX.parent.mkdir(parents=True, exist_ok=True)
        wb.save(NOTICE_BOARD_XLSX)
        return str(NOTICE_BOARD_XLSX)
    except Exception:
        return ""


def _broadcast_notice_to_citizens(record: dict) -> int:
    """Push a published notice to every active citizen's inbox (notifications)."""
    if not record or record.get("status") != "Published":
        return 0
    now = datetime.now(timezone.utc).isoformat()
    title = record.get("title") or "Gram Panchayat Notice"
    message = record.get("description") or title
    if len(message) > 400:
        message = message[:397] + "..."

    import uuid

    user_ids = []
    offset = 0
    while True:
        page = (
            supabase.table("users")
            .select("user_id")
            .eq("role", "villager")
            .eq("account_status", "active")
            .order("user_id")
            .range(offset, offset + 999)
            .execute()
            .data
            or []
        )
        user_ids.extend(r["user_id"] for r in page)
        if len(page) < 1000:
            break
        offset += 1000

    sent = 0
    for i in range(0, len(user_ids), 500):
        chunk = user_ids[i:i + 500]
        payload = [
            {
                "notification_id": str(uuid.uuid4()),
                "user_id": uid,
                # "Announcement" is the only value accepted by the
                # notifications_notification_type_check constraint.
                "notification_type": "Announcement",
                "title": title,
                "message": message,
                "reference_type": "announcement",
                "reference_id": record.get("announcement_id"),
                "is_read": False,
                "created_at": now,
            }
            for uid in chunk
        ]
        try:
            supabase.table("notifications").insert(payload).execute()
            sent += len(payload)
        except Exception as exc:
            # A single bad row (e.g. user deleted mid-broadcast) fails the whole
            # batch — retry row-by-row so the rest of the chunk still lands.
            print(f"[notice-broadcast] chunk insert failed, retrying row-wise: {exc}", flush=True)
            for row in payload:
                try:
                    supabase.table("notifications").insert(row).execute()
                    sent += 1
                except Exception as row_exc:
                    print(f"[notice-broadcast] skipped {row['user_id']}: {row_exc}", flush=True)
            continue
    return sent


def _purge_notice_notifications(announcement_id: str):
    """Drop inbox copies of a notice once the notice itself is deleted."""
    try:
        supabase.table("notifications").delete().eq("reference_id", announcement_id).execute()
    except Exception as exc:
        print(f"[notice-broadcast] inbox purge failed: {exc}", flush=True)


# ---- Scheduled publishing ---------------------------------------------------
# Notices can be Published with a FUTURE published_at. Citizens only see them
# (and only get the inbox broadcast) once that timestamp passes. A throttled
# sweep on the notice list picks up notices whose schedule has just elapsed —
# restart-safe because it skips any notice that already has inbox messages.
_LAST_SWEEP_AT = 0.0
_SWEEP_INTERVAL_SECONDS = 60
_SWEEP_LOOKBACK_DAYS = 3


def _maybe_schedule_sweep(background: BackgroundTasks):
    global _LAST_SWEEP_AT
    if time.time() - _LAST_SWEEP_AT < _SWEEP_INTERVAL_SECONDS:
        return
    _LAST_SWEEP_AT = time.time()
    background.add_task(_sweep_scheduled_notices)


def _sweep_scheduled_notices():
    now = datetime.now(timezone.utc)
    since = (now - timedelta(days=_SWEEP_LOOKBACK_DAYS)).isoformat()
    now_iso = now.isoformat()
    try:
        rows = (
            supabase.table("announcements")
            .select("*")
            .eq("village_id", KUMBHARI_VILLAGE_ID)
            .eq("status", "Published")
            .gte("published_at", since)
            .lte("published_at", now_iso)
            .execute()
            .data
            or []
        )
    except Exception as exc:
        print(f"[scheduled-notice] sweep query failed: {exc}", flush=True)
        return

    for record in rows:
        try:
            existing = (
                supabase.table("notifications")
                .select("notification_id")
                .eq("reference_id", record.get("announcement_id"))
                .limit(1)
                .execute()
            )
            if existing.data:
                continue  # already announced when it went live
            sent = _broadcast_notice_to_citizens(record)
            print(f"[scheduled-notice] swept {record.get('announcement_id')} -> {sent} inboxes", flush=True)
        except Exception as exc:
            print(f"[scheduled-notice] sweep failed for {record.get('announcement_id')}: {exc}", flush=True)


_MISSING_COL_RE = re.compile(r"Could not find the '([^']+)' column")


def _write_notice(action, payload):
    """Insert/update a notice, tolerating schema drift.

    If PostgREST reports a column that does not exist (PGRST204) — e.g. the
    optional ``event_date`` column on an older database — drop that column and
    retry so the desk is never blocked by a missing column.
    """
    from postgrest.exceptions import APIError

    data = dict(payload)
    for _ in range(4):
        try:
            return action(data).execute()
        except APIError as exc:
            match = _MISSING_COL_RE.search(str(exc))
            if match and match.group(1) in data:
                data.pop(match.group(1), None)
                continue
            raise
    return action(data).execute()


def _is_live(record: dict) -> bool:
    """True when a Published notice's scheduled publish time has passed (or none)."""
    pub = record.get("published_at")
    if not pub:
        return True
    return str(pub) <= datetime.now(timezone.utc).isoformat()


def _notice_sort_key(record: dict) -> str:
    """Board ordering date: the notice's own date, else publish/creation time.
    event_date values are YYYY-MM-DD and timestamps are ISO-8601, so the
    strings sort chronologically."""
    return str(
        record.get("event_date")
        or record.get("published_at")
        or record.get("created_at")
        or ""
    )


def _sort_notices(records: list) -> list:
    """Newest notice first by board date; notices without a date go last."""
    return sorted(records, key=_notice_sort_key, reverse=True)


def _on_notice_change(background: BackgroundTasks, broadcast: Optional[dict] = None, purge: Optional[str] = None):
    """Real-time side effects: refresh the Excel sheet, notify citizens."""
    background.add_task(_sync_notice_excel)
    if broadcast:
        background.add_task(_broadcast_notice_to_citizens, broadcast)
    if purge:
        background.add_task(_purge_notice_notifications, purge)


@router.get("")
def list_announcements(
    background: BackgroundTasks,
    current_user: CurrentUser = Depends(get_current_user),
):
    query = supabase.table("announcements").select("*")
    query = query.eq("village_id", KUMBHARI_VILLAGE_ID)
    _maybe_schedule_sweep(background)
    if current_user.role == "villager":
        # Villagers only ever see published, live (not scheduled-away), non-expired
        # notices, and the board keeps only the newest CITIZEN_NOTICE_LIMIT.
        today = datetime.now(timezone.utc).date().isoformat()
        now_iso = datetime.now(timezone.utc).isoformat()
        query = query.eq("status", "Published").or_(
            f"expiry_date.is.null,expiry_date.gte.{today}"
        )
        # published_at <= now also excludes NULLs, so fetch a few extra rows and
        # keep only the ones that are actually live before slicing to the limit.
        result = query.order("created_at", desc=True).limit(CITIZEN_NOTICE_LIMIT * 4).execute()
        live = [a for a in (result.data or []) if _is_live(a)]
        return _sort_notices(live)[:CITIZEN_NOTICE_LIMIT]
    result = query.order("created_at", desc=True).limit(STAFF_NOTICE_LIMIT).execute()
    return _sort_notices(result.data or [])


@router.get("/board-excel")
def download_notice_board_excel(current_user: CurrentUser = Depends(get_current_user)):
    """Download the live Digital Notice Board Excel sheet."""
    if not NOTICE_BOARD_XLSX.exists():
        _sync_notice_excel()
    if not NOTICE_BOARD_XLSX.exists():
        raise HTTPException(status_code=404, detail="Excel sheet not generated yet")
    from fastapi.responses import FileResponse

    return FileResponse(
        str(NOTICE_BOARD_XLSX),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="kumbhari_notice_board.xlsx",
    )


@router.post("", status_code=201)
def create_announcement(
    payload: AnnouncementCreate,
    background: BackgroundTasks,
    request: Request,
    current_user: CurrentUser = Depends(require_roles("panchayat_staff", "panchayat_admin", "system_admin")),
):
    raw = payload.model_dump(exclude_unset=True)
    raw.pop("village_id", None)
    raw["village_id"] = KUMBHARI_VILLAGE_ID
    inserted = _write_notice(
        lambda data: supabase.table("announcements").insert(data),
        {**raw, "created_by": current_user.user_id, "status": "Draft"},
    )
    record = inserted.data[0]
    log_audit(
        actor_id=current_user.user_id,
        action="Create Notice",
        entity_type="announcement",
        entity_id=record["announcement_id"],
        new_values={"title": payload.title, "status": "Draft"},
        request=request,
    )
    _on_notice_change(background)
    return record


@router.patch("/{announcement_id}")
def update_announcement(
    announcement_id: str,
    updates: AnnouncementUpdate,
    background: BackgroundTasks,
    request: Request,
    current_user: CurrentUser = Depends(require_roles("panchayat_staff", "panchayat_admin", "system_admin")),
):
    data = updates.model_dump(exclude_unset=True)
    data.pop("village_id", None)
    publish_at = data.pop("publish_at", None)  # not a DB column; mapped to published_at below
    if publish_at:
        try:
            dt = datetime.fromisoformat(str(publish_at).strip().replace("Z", "+00:00"))
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            publish_at = dt.astimezone(timezone.utc).isoformat()
        except ValueError:
            raise HTTPException(status_code=400, detail="publish_at must be an ISO-8601 timestamp")
    if not data and publish_at is None:
        raise HTTPException(status_code=400, detail="Nothing to update")
    if data.get("status") and data["status"] not in VALID_STATUSES:
        raise HTTPException(status_code=400, detail="Status must be one of: Draft, Published, Archived")

    prev = supabase.table("announcements").select("status").eq("announcement_id", announcement_id).execute()
    if not prev.data:
        raise HTTPException(status_code=404, detail="Notice not found")
    was_published = prev.data[0].get("status") == "Published"

    if data.get("status") == "Published":
        # publish_at (optional) schedules the release; otherwise publish now.
        data["published_at"] = publish_at or datetime.now(timezone.utc).isoformat()
        data["reviewed_by"] = current_user.user_id
    elif data.get("status") == "Archived" or (was_published and data.get("status") == "Draft"):
        data["published_at"] = None

    if not data:
        raise HTTPException(status_code=400, detail="Nothing to update")

    result = _write_notice(
        lambda data: supabase.table("announcements").update(data).eq("announcement_id", announcement_id),
        data,
    )
    if not result.data:
        raise HTTPException(status_code=404, detail="Notice not found")

    scheduled = data.get("status") == "Published" and data.get("published_at", "") > datetime.now(timezone.utc).isoformat()
    action = "Schedule Notice" if scheduled else ("Publish Notice" if data.get("status") == "Published" else "Edit Notice")
    log_audit(
        actor_id=current_user.user_id,
        action=action,
        entity_type="announcement",
        entity_id=announcement_id,
        old_values=prev.data[0],
        new_values={**data, **({"publish_at": publish_at} if publish_at else {})},
        request=request,
    )

    became_published = data.get("status") == "Published" and not was_published
    # Only notify citizens when the notice is actually live right now;
    # scheduled notices are picked up by the sweep once publish time passes.
    broadcast_record = result.data[0] if became_published and _is_live(result.data[0]) else None
    _on_notice_change(background, broadcast=broadcast_record)
    return result.data[0]


@router.delete("/{announcement_id}", status_code=200)
def delete_announcement(
    announcement_id: str,
    background: BackgroundTasks,
    request: Request,
    current_user: CurrentUser = Depends(require_roles("panchayat_staff", "panchayat_admin", "system_admin")),
):
    prev = supabase.table("announcements").select("*").eq("announcement_id", announcement_id).execute()
    if not prev.data:
        raise HTTPException(status_code=404, detail="Notice not found")

    supabase.table("announcements").delete().eq("announcement_id", announcement_id).execute()
    log_audit(
        actor_id=current_user.user_id,
        action="Delete Notice",
        entity_type="announcement",
        entity_id=announcement_id,
        new_values={"deleted_title": prev.data[0].get("title")},
        request=request,
    )
    _on_notice_change(background, purge=announcement_id)
    return {"deleted": announcement_id, "title": prev.data[0].get("title")}